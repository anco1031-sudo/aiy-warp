#!/usr/bin/env python3
"""
office/sheet_write.py — Excel generation from JSON (Aetox 4.2)
Usage: python3 scripts/office/sheet_write.py <input.json> [--out report.xlsx]
JSON schema: {"path": "report.xlsx", "sheets": [{"name":"Sales","columns":[],"formats":[],"rows":[]}]}
Deps: openpyxl (pip install openpyxl)
"""
import json, sys, os, argparse
def write_sheet(data: dict, out_path: str=None) -> str:
    try: import openpyxl
    except ImportError: return "ERROR: pip install openpyxl"
    from openpyxl import Workbook
    path = out_path or data.get("path","report.xlsx")
    wb=Workbook()
    for i, sh in enumerate(data.get("sheets",[])):
        ws = wb.active if i==0 else wb.create_sheet(sh["name"])
        if i==0: ws.title=sh["name"]
        for col, h in enumerate(sh.get("columns",[]),1): ws.cell(row=1,column=col,value=h).font=openpyxl.styles.Font(bold=True)
        for r_idx, row in enumerate(sh.get("rows",[]),2):
            for c_idx, v in enumerate(row,1): ws.cell(row=r_idx,column=c_idx,value=v)
        for c_idx, fmt in enumerate(sh.get("formats",[]),1):
            if fmt=="money":
                for r in range(2, len(sh.get("rows",[]))+2): ws.cell(row=r,column=c_idx).number_format='#,##0.00'
            elif fmt=="percent":
                for r in range(2, len(sh.get("rows",[]))+2): ws.cell(row=r,column=c_idx).number_format='0.00%'
    wb.save(path); return path

if __name__=="__main__":
    ap=argparse.ArgumentParser()
    ap.add_argument("json_file", help="JSON file or '-' for stdin")
    ap.add_argument("--out", default=None)
    args=ap.parse_args()
    if args.json_file=="-": data=json.load(sys.stdin)
    else:
        with open(args.json_file) as f: data=json.load(f)
    res=write_sheet(data, args.out)
    print(res)
